from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "NCPOR_Polar_Portal_Project_Report.docx"

NAVY = "08263A"
PALE_BLUE = "EAF3F7"
MID_BLUE = "D5E7EF"
GRAY = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def set_cell_border(cell, color=GRAY, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:color"), color)


def set_cell_margins(cell, top=150, start=180, bottom=150, end=180):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_row_height(row, twips):
    tr_pr = row._tr.get_or_add_trPr()
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), str(twips))
    height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(height)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(15)
        style.paragraph_format.space_after = Pt(7)


def add_paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        lead.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_screenshot_placeholder(doc, caption, label):
    caption_p = doc.add_paragraph()
    caption_p.paragraph_format.space_before = Pt(7)
    caption_p.paragraph_format.space_after = Pt(4)
    run = caption_p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)

    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cell, PALE_BLUE)
    set_cell_border(cell, color="9ABAC9", size="10")
    set_cell_margins(cell, top=260, start=260, bottom=260, end=260)
    set_row_height(table.rows[0], 3100)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(39, 90, 112)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Paste or insert the relevant screenshot in this area")
    r2.italic = True
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(79, 106, 120)


def add_results_table(doc):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(5.3), Cm(10.2)]
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, title in enumerate(("Area", "Observed outcome")):
        cell = header.cells[index]
        cell.width = widths[index]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    rows = [
        ("Public discovery", "A responsive public shell provides access to the home page, explorer, expeditions, knowledge archive, stories and education pathways."),
        ("Knowledge archive", "The archive presents database-backed expedition records with search, type navigation, filters, pagination and direct record routes."),
        ("Cloud data path", "Supabase is structured as the cloud source of truth for PostgreSQL data, authentication, storage and role-aware access controls."),
        ("Administration", "Protected administration routes provide archive management, record creation and cover-image support through the existing Supabase architecture."),
        ("Outreach workflow", "The Outreach Studio is designed to generate source-grounded draft content through a protected server-side AI integration."),
    ]
    for row_index, (area, outcome) in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate((area, outcome)):
            cell = row.cells[index]
            cell.width = widths[index]
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            run.font.size = Pt(9.5)
            if index == 0:
                run.bold = True
    doc.add_paragraph()


def add_reference(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    p.add_run(f"{number}. ").bold = True
    p.add_run(text)


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    configure_styles(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("NCPOR POLAR KNOWLEDGE AND OUTREACH PORTAL")
    header_run.font.size = Pt(8.5)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor(48, 71, 85)
    add_page_number(section.footer.paragraphs[0])

    # Cover page
    doc.add_paragraph().paragraph_format.space_after = Pt(26)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("NCPOR Polar Knowledge and Outreach Portal")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(8)
    subtitle.paragraph_format.space_after = Pt(22)
    run = subtitle.add_run("Project Report")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(32, 76, 96)
    details = [
        "Problem Statement: Integrated Polar Science Outreach, Knowledge Repository and Media Dissemination Portal",
        "Organization: National Centre for Polar and Ocean Research",
        "Theme: Smart Education",
        "Prepared for project presentation",
    ]
    for item in details:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        p.add_run(item).font.size = Pt(11)
    doc.add_paragraph().paragraph_format.space_after = Pt(105)
    submitted = doc.add_paragraph()
    submitted.alignment = WD_ALIGN_PARAGRAPH.CENTER
    submitted.add_run("Submitted by: ______________________________").font.size = Pt(11)
    doc.add_page_break()

    add_heading(doc, "Abstract")
    add_paragraph(doc, "The NCPOR Polar Knowledge and Outreach Portal is a web-based platform developed to make polar-science information easier to discover, understand and reuse. The project responds to the need for a connected digital environment where expedition information, research themes, publications, datasets, media and outreach material can be accessed through one coherent interface. Rather than presenting scientific records as isolated files or disconnected lists, the portal is designed around relationships between people, places, expeditions and knowledge objects.")
    add_paragraph(doc, "The implemented solution combines a modern Next.js application with Supabase Cloud services for PostgreSQL data, authentication, storage and role-aware access. The public side emphasizes an editorial polar visual language, responsive navigation and a unified archive, while the administrative side provides authenticated workflows for creating and managing records. An Outreach Studio extends the platform by preparing audience-specific communication drafts from selected source content through a protected AI integration. The system uses synthetic development content where institutional records are not available, and the interface identifies this status clearly to avoid presenting demonstrations as official scientific findings.")
    add_paragraph(doc, "The resulting portal establishes a practical foundation for a future NCPOR knowledge ecosystem. It supports public discovery, structured content management and cloud-backed growth without requiring a separate local production database.")

    add_heading(doc, "1. Introduction")
    add_paragraph(doc, "Polar research generates a wide range of valuable information, including expedition records, observations, publications, datasets, images, reports and public outreach material. These resources are often difficult for non-specialist audiences to discover and can be fragmented across different systems and formats. The NCPOR Polar Knowledge and Outreach Portal was conceived as an integrated response to this challenge.")
    add_paragraph(doc, "The project aims to create a credible digital field station for knowledge: a space that remains scientifically structured while being approachable for students, educators, researchers, outreach teams and the general public. It combines a discovery-oriented public experience with a research-oriented archive so that users can move from a broad topic or visual story to detailed records and related material without losing context.")
    add_paragraph(doc, "The current implementation focuses on a strong vertical slice of this vision. It includes a polished public portal, a unified archive, Supabase-backed content access, authentication-aware administration and an outreach workflow. The architecture is deliberately modular so that additional content types and advanced exploration features can be added without replacing the core application.")

    add_heading(doc, "2. Methodology")
    add_heading(doc, "2.1 Requirements and experience design", level=2)
    add_paragraph(doc, "The work began by translating the product, architecture, user-experience and design specifications into a usable application foundation. The design direction combines restrained polar editorial styling with scientific clarity. Priority was given to information hierarchy, readable metadata, responsive behavior, keyboard-accessible controls and reduced-motion support.")
    add_heading(doc, "2.2 Application architecture", level=2)
    add_paragraph(doc, "The portal was developed as a Next.js and TypeScript web application. Reusable components were organized around the public shell, archive discovery, administration and outreach workflows. The user interface communicates with repository and service layers rather than assembling complex database queries directly in page components. This keeps the data path easier to maintain and test.")
    add_heading(doc, "2.3 Cloud backend integration", level=2)
    add_paragraph(doc, "Supabase Cloud was selected as the production source of truth. The platform uses Supabase PostgreSQL for structured content, Supabase Auth for sessions and administrator access, and Supabase Storage for controlled media and cover-image handling. Browser and server clients are separated so that the publishable key can support approved public/client operations while privileged credentials remain server-only. Database migration and seed files provide a repeatable setup for the project schema and synthetic demonstration records.")
    add_heading(doc, "2.4 Verification approach", level=2)
    add_paragraph(doc, "The project was checked through TypeScript type validation, ESLint, unit tests, Playwright browser tests and a Next.js production build. The archive flow was inspected using database-backed public records, direct query URLs and detail routes. Environment configuration was also documented for Vercel deployment because public Next.js variables must be available during the build as well as in the deployed runtime.")

    add_heading(doc, "3. Observations")
    add_paragraph(doc, "The public homepage demonstrates that a formal research portal can use a contemporary visual language without losing institutional credibility. The navigation remains concise, while the content hierarchy directs users from introductory material toward discovery routes such as expeditions, the knowledge archive, stories and education.")
    add_paragraph(doc, "The unified archive is the most visible example of the knowledge-first approach. It presents records as discoverable objects rather than as a conventional database table. Each result combines a type label, title, concise description, temporal context, regional context and selected metadata. Search queries and filters are designed to be represented in the URL so that a selected discovery state can be reopened or shared.")
    add_paragraph(doc, "The administrative workspace shows the importance of connecting editorial actions to the underlying database. Archive records created through the management interface are intended to follow the same Supabase-backed data path as public records. Support for record types, publication status, metadata fields and cover images provides a practical base for more complete editorial workflows.")
    add_paragraph(doc, "The deployment review highlighted that cloud architecture depends on configuration as well as code. The public Supabase URL and publishable key must be configured in Vercel as NEXT_PUBLIC variables so that the browser and server clients can initialize correctly. Server-only keys, where required, must remain outside client bundles.")

    # Required screenshot space immediately after observations.
    doc.add_page_break()
    add_heading(doc, "Screenshots and Interface Evidence")
    add_paragraph(doc, "Insert presentation screenshots in the spaces below after capturing the final deployed portal.")
    add_screenshot_placeholder(doc, "Figure 1. Public knowledge archive or homepage", "SCREENSHOT SPACE 1")
    add_screenshot_placeholder(doc, "Figure 2. Administrative archive management or Outreach Studio", "SCREENSHOT SPACE 2")
    doc.add_page_break()

    add_heading(doc, "4. Results")
    add_paragraph(doc, "The project produced a functioning web portal foundation that connects public discovery with cloud-backed content management. The implemented archive and detail experiences use a normalized data-access path rather than static frontend arrays, and the user interface provides meaningful loading, empty and error states when content is unavailable.")
    add_results_table(doc)
    add_paragraph(doc, "The project also establishes a deployment-ready environment contract. The browser-safe Supabase URL and publishable key use NEXT_PUBLIC names, while privileged and AI credentials remain server-only. This separation supports secure deployment on Vercel without embedding privileged database access in the browser.")

    add_heading(doc, "5. Conclusion")
    add_paragraph(doc, "The NCPOR Polar Knowledge and Outreach Portal demonstrates how a scientific institution can present polar research through a connected, modern and extensible digital platform. Its central contribution is not only visual presentation; it is the combination of discoverability, metadata, cloud-backed records and editorial control within one application structure.")
    add_paragraph(doc, "The current product slice provides the core ingredients for continued expansion: a responsive public portal, a unified archive, record detail routes, protected administrative workflows, Supabase integration and an AI-assisted outreach pathway. The architecture supports future work on richer publications, datasets, maps, media galleries, education resources and relationship-driven exploration. With verified cloud configuration and continued curation of authoritative NCPOR content, the platform can evolve from a synthetic demonstration into a sustainable institutional knowledge and outreach system.")

    add_heading(doc, "6. References")
    references = [
        "NCPOR Polar Knowledge and Outreach Portal. Master Product Specification. Project specification pack.",
        "NCPOR Polar Knowledge and Outreach Portal. Architecture Blueprint. Project specification pack.",
        "NCPOR Polar Knowledge and Outreach Portal. Data Model Requirements. Project specification pack.",
        "NCPOR Polar Knowledge and Outreach Portal. User Experience Specification and Design System Specification. Project specification pack.",
        "Supabase. API Keys and Next.js environment configuration documentation. https://supabase.com/docs/guides/getting-started/api-keys",
        "Vercel. Environment Variables documentation. https://vercel.com/docs/environment-variables",
        "Next.js. Documentation for the App Router and environment variables. https://nextjs.org/docs",
    ]
    for number, reference in enumerate(references, start=1):
        add_reference(doc, number, reference)

    doc.core_properties.title = "NCPOR Polar Knowledge and Outreach Portal Project Report"
    doc.core_properties.subject = "Project report"
    doc.core_properties.author = "NCPOR Polar Portal Project Team"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
